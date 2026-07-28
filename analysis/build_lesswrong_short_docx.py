#!/usr/bin/env python3
"""Build the short LessWrong article as a polished Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "lesswrong_article_short.md"
OUTPUT = ROOT / "docs" / "lesswrong_article_short.docx"

FONT = "Aptos"
INK = RGBColor(31, 55, 72)
BODY = RGBColor(28, 28, 28)
MUTED = RGBColor(92, 92, 92)
ACCENT = RGBColor(177, 91, 48)
LINK = RGBColor(34, 93, 152)


def set_run_font(
    run,
    *,
    name: str = FONT,
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def add_hyperlink(paragraph, text: str, url: str, *, italic: bool = False):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), f"{LINK[0]:02X}{LINK[1]:02X}{LINK[2]:02X}")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "21")
    r_pr.extend([fonts, color, underline, size, size_cs])
    if italic:
        r_pr.append(OxmlElement("w:i"))
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\*([^*]+)\*)"
)


def add_inline(paragraph, text: str, *, default_italic: bool = False):
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=10.5, color=BODY, italic=default_italic)
        token = match.group(0)
        if token.startswith("["):
            label, url = match.group(2), match.group(3)
            italic = label.startswith("*") and label.endswith("*")
            add_hyperlink(paragraph, label.strip("*"), url, italic=italic)
        elif token.startswith("**"):
            run = paragraph.add_run(match.group(4))
            set_run_font(run, size=10.5, color=BODY, bold=True, italic=default_italic)
        elif token.startswith("`"):
            run = paragraph.add_run(match.group(5))
            set_run_font(
                run,
                name="Aptos Mono",
                size=9.6,
                color=RGBColor(70, 70, 70),
                italic=default_italic,
            )
        else:
            run = paragraph.add_run(match.group(6))
            set_run_font(run, size=10.5, color=BODY, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=10.5, color=BODY, italic=default_italic)


def add_custom_bullet_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(el.get(qn("w:abstractNumId")))
        for el in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    level.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def set_bullet(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def set_image_alt_text(paragraph, title: str, description: str):
    drawings = paragraph._p.xpath(".//wp:docPr")
    if not drawings:
        return
    drawings[0].set("title", title)
    drawings[0].set("descr", description)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.widow_control = True

    for style_name, size, before, after in (
        ("Heading 1", 15.5, 12, 5),
        ("Heading 2", 12.5, 9, 4),
        ("Heading 3", 11.5, 7, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Figure Caption" not in [style.name for style in doc.styles]:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Figure Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.line_spacing = 1.05
    caption.paragraph_format.keep_together = True

    if "Pull Quote" not in [style.name for style in doc.styles]:
        quote = doc.styles.add_style("Pull Quote", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = doc.styles["Pull Quote"]
    quote.font.name = FONT
    quote._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    quote._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    quote.font.size = Pt(10.5)
    quote.font.italic = True
    quote.font.color.rgb = INK
    quote.paragraph_format.left_indent = Inches(0.35)
    quote.paragraph_format.right_indent = Inches(0.25)
    quote.paragraph_format.space_before = Pt(3)
    quote.paragraph_format.space_after = Pt(7)


def add_left_border(paragraph, color: str = "B15B30", size: str = "16"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)


def parse_blocks(markdown: str):
    lines = markdown.splitlines()
    blocks = []
    paragraph_lines = []

    def flush():
        if paragraph_lines:
            blocks.append(("paragraph", " ".join(line.strip() for line in paragraph_lines)))
            paragraph_lines.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("!["):
            flush()
            match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", stripped)
            blocks.append(("image", match.groups()))
        elif stripped.startswith("# "):
            flush()
            blocks.append(("title", stripped[2:]))
        elif stripped.startswith("## "):
            flush()
            blocks.append(("heading", stripped[3:]))
        elif stripped.startswith("> "):
            flush()
            blocks.append(("quote", stripped[2:]))
        elif stripped.startswith("- "):
            flush()
            paragraph_lines.append(stripped[2:])
            # Wrapped list lines are folded into the same bullet below.
            blocks.append(("bullet_pending", None))
        elif blocks and blocks[-1][0] == "bullet_pending" and paragraph_lines:
            paragraph_lines.append(stripped)
        elif stripped.startswith("*Figure ") and stripped.endswith("*"):
            flush()
            blocks.append(("caption", stripped[1:-1]))
        elif (
            stripped.startswith("*")
            and stripped.endswith("*")
            and not stripped.startswith("**")
        ):
            flush()
            blocks.append(("subtitle", stripped[1:-1]))
        else:
            if blocks and blocks[-1][0] == "bullet_pending":
                paragraph_lines.append(stripped)
            else:
                paragraph_lines.append(stripped)
        if blocks and blocks[-1][0] == "bullet_pending" and paragraph_lines:
            # A bullet ends when its source line has no continuation marker.
            # Continuations are handled by the indentation test in the second pass.
            pass
    flush()

    # Reparse bullets with a simpler line-oriented pass to preserve continuations.
    normalized = []
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if stripped.startswith("- "):
            content = [stripped[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("  "):
                content.append(lines[i].strip())
                i += 1
            normalized.append(("bullet", " ".join(content)))
            continue
        i += 1

    # Replace provisional bullet material using a clean full pass.
    result = []
    i = 0
    paragraph_lines = []

    def flush_result():
        if paragraph_lines:
            result.append(("paragraph", " ".join(x.strip() for x in paragraph_lines)))
            paragraph_lines.clear()

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not stripped:
            flush_result()
        elif stripped.startswith("- "):
            flush_result()
            content = [stripped[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("  "):
                content.append(lines[i].strip())
                i += 1
            result.append(("bullet", " ".join(content)))
            continue
        elif stripped.startswith("!["):
            flush_result()
            match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", stripped)
            result.append(("image", match.groups()))
        elif stripped.startswith("# "):
            flush_result()
            result.append(("title", stripped[2:]))
        elif stripped.startswith("## "):
            flush_result()
            result.append(("heading", stripped[3:]))
        elif stripped.startswith("> "):
            flush_result()
            result.append(("quote", stripped[2:]))
        elif stripped.startswith("*Figure ") and stripped.endswith("*"):
            flush_result()
            result.append(("caption", stripped[1:-1]))
        elif stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            flush_result()
            result.append(("subtitle", stripped[1:-1]))
        else:
            paragraph_lines.append(stripped)
        i += 1
    flush_result()
    return result


def build():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    # Named compact-editorial override to narrative_proposal.
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)

    configure_styles(doc)
    bullet_num_id = add_custom_bullet_numbering(doc)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("INOCULATE OR REFLECT?")
    set_run_font(run, size=7.5, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    add_page_number(footer_p)

    image_count = 0
    first_body = True
    for kind, content in parse_blocks(SOURCE.read_text()):
        if kind == "title":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(content)
            set_run_font(run, size=27, color=INK, bold=True)
        elif kind == "subtitle":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(content)
            set_run_font(run, size=12.5, color=MUTED, italic=True)
        elif kind == "heading":
            doc.add_paragraph(content, style="Heading 1")
        elif kind == "quote":
            p = doc.add_paragraph(style="Pull Quote")
            add_left_border(p)
            add_inline(p, content, default_italic=True)
        elif kind == "bullet":
            p = doc.add_paragraph()
            set_bullet(p, bullet_num_id)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, content)
        elif kind == "image":
            alt, path_string = content
            path = Path(path_string)
            image_count += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            width = Inches(6.7 if image_count == 1 else 5.45)
            p.add_run().add_picture(str(path), width=width)
            set_image_alt_text(p, f"Figure {image_count}", alt)
        elif kind == "caption":
            p = doc.add_paragraph(style="Figure Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, content, default_italic=True)
            for run in p.runs:
                set_run_font(run, size=8.5, color=MUTED, italic=True)
        else:
            p = doc.add_paragraph()
            if first_body:
                p.paragraph_format.space_before = Pt(2)
                first_body = False
            add_inline(p, content)

    final_link = doc.paragraphs[-1]
    final_link.paragraph_format.space_before = Pt(5)
    final_link.paragraph_format.keep_together = True

    props = doc.core_properties
    props.title = "Inoculate or Reflect?"
    props.subject = "Two training interventions under prompting, steering, and patching"
    props.author = ""
    props.keywords = "Inoculation Prompting, Counterfactual Reflection Training, sycophancy"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
